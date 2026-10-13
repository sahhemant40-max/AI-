import os
import uuid
from datetime import datetime, timezone
from typing import Dict, Any, List
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from backend.app.core.config import settings
from backend.app.database.session import get_database
from backend.app.services.scan_service import get_scan_by_id
from backend.app.services.model_service import get_model_by_id
from backend.app.services.finding_service import list_findings
from backend.app.core.logging import logger
from backend.app.core.errors import ResourceNotFoundException


def generate_pdf_report(report_data: Dict[str, Any], filepath: str):
    """Compiles professional cybersecurity PDF report using ReportLab."""
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    styles = getSampleStyleSheet()

    # Custom Cyber Styles
    title_style = ParagraphStyle(
        'CyberTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        textColor=colors.HexColor('#0E1D15'),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        'CyberSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#19C96B'),
        spaceAfter=15,
    )
    heading_style = ParagraphStyle(
        'CyberHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        textColor=colors.HexColor('#0E1D15'),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        'CyberBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        textColor=colors.HexColor('#333333'),
        leading=13,
    )

    story = []

    # Header / Branding
    story.append(Paragraph("AI SENTINEL — SECURITY & RISK ASSESSMENT REPORT", title_style))
    story.append(Paragraph("Tagline: Secure AI. Measure Risk. Build Trust. | Enterprise AI Governance Dossier", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#19C96B'), spaceAfter=15))

    # Meta Table
    scan = report_data["scan"]
    model = report_data["model"]

    meta_table_data = [
        ["Report ID:", report_data["id"], "Generated At:", report_data["created_at"][:19]],
        ["Target Model:", model.get("name", "N/A"), "Model Architecture:", model.get("model_type", "N/A")],
        ["Framework:", model.get("framework", "N/A"), "Environment:", model.get("environment", "N/A")],
        ["Scan ID:", scan["id"], "Scan Intensity:", scan.get("intensity", "comprehensive").upper()],
        ["Overall Risk Score:", f"{scan.get('risk_score', 'N/A')} / 100", "Risk Tier:", scan.get("risk_category", "N/A").upper()],
    ]
    t_meta = Table(meta_table_data, colWidths=[110, 160, 110, 160])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F4FBF6')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#0E1D15')),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1E7DD')),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 15))

    # Section 1: Executive Summary
    story.append(Paragraph("1. EXECUTIVE ASSESSMENT SUMMARY", heading_style))
    score = scan.get('risk_score', 0)
    category = scan.get('risk_category', 'Low')
    summary_text = (
        f"This automated audit evaluated the '{model.get('name')}' model across requested security vectors. "
        f"The composite deterministic risk score is <b>{score}/100 ({category} Risk)</b> with an assessment completeness "
        f"rate of <b>{int((scan.get('completeness') or 1.0) * 100)}%</b>. "
        f"A total of <b>{len(report_data['findings'])} security findings</b> were identified."
    )
    story.append(Paragraph(summary_text, body_style))
    story.append(Spacer(1, 12))

    # Section 2: Security Findings
    story.append(Paragraph("2. VULNERABILITY & BIAS FINDINGS", heading_style))
    findings = report_data["findings"]
    if not findings:
        story.append(Paragraph("No critical vulnerabilities or severe bias disparities detected.", body_style))
    else:
        find_table = [["ID", "Severity", "Module", "Finding Title", "Status"]]
        for f in findings:
            find_table.append([
                f["id"],
                f.get("severity", "Medium"),
                f.get("module", "Security"),
                f.get("title", ""),
                f.get("status", "open").upper()
            ])
        t_find = Table(find_table, colWidths=[65, 60, 110, 240, 65])
        t_find.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0E1D15')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_find)

    story.append(Spacer(1, 15))

    # Section 3: NIST AI RMF Governance Mapping
    story.append(Paragraph("3. NIST AI RMF 1.0 GOVERNANCE MAPPING", heading_style))
    nist_text = (
        "Empirical assessment evidence generated during this scan cross-maps to NIST AI RMF core functions: "
        "GOVERN (Accountability & risk policies), MAP (Context & categorization), "
        "MEASURE (Quantitative robustness, fairness, and privacy metrics), and MANAGE (Remediation lifecycle)."
    )
    story.append(Paragraph(nist_text, body_style))
    story.append(Spacer(1, 15))

    # Disclaimer
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#999999'), spaceAfter=8))
    story.append(Paragraph("AI Sentinel &bull; Defensive AI Security & Governance Platform &bull; Automated Telemetry Report", body_style))

    doc.build(story)


def generate_docx_report(report_data: Dict[str, Any], filepath: str):
    """Compiles Word DOCX report using python-docx."""
    doc = Document()
    doc.add_heading("AI SENTINEL — AI Model Risk Assessment Report", level=0)

    scan = report_data["scan"]
    model = report_data["model"]

    p = doc.add_paragraph()
    p.add_run("Tagline: ").bold = True
    p.add_run("Secure AI. Measure Risk. Build Trust.\n")
    p.add_run(f"Report ID: {report_data['id']}\n")
    p.add_run(f"Generated: {report_data['created_at']}\n")
    p.add_run(f"Target Model: {model.get('name')} (v{model.get('version')})\n")
    p.add_run(f"Risk Score: {scan.get('risk_score', 'N/A')}/100 ({scan.get('risk_category', 'N/A')})\n")

    doc.add_heading("1. Assessment Overview", level=1)
    doc.add_paragraph(
        f"The AI Sentinel engine executed security probes against '{model.get('name')}'. "
        f"Scan ID: {scan['id']}. Tested modules: {', '.join(scan.get('modules', []))}."
    )

    doc.add_heading("2. Findings Summary", level=1)
    for f in report_data["findings"]:
        p_find = doc.add_paragraph()
        p_find.add_run(f"[{f.get('severity')}] {f.get('title')}\n").bold = True
        p_find.add_run(f"Module: {f.get('module')} | ID: {f.get('id')}\n")
        p_find.add_run(f"Description: {f.get('description')}\n")
        p_find.add_run(f"Remediation: {f.get('remediation', 'N/A')}\n")

    doc.add_heading("3. NIST AI RMF Alignment", level=1)
    doc.add_paragraph(
        "Technical evidence aligns with NIST AI RMF 1.0 guidelines across Govern, Map, Measure, and Manage functions."
    )

    doc.save(filepath)


async def create_report(scan_id: str, report_type: str = "Comprehensive Technical Audit", title: str = "AI Sentinel Risk Assessment Report") -> Dict[str, Any]:
    """Generates PDF and DOCX reports from saved scan data and stores metadata in MongoDB."""
    db = get_database()
    scan = await get_scan_by_id(scan_id)
    model = await get_model_by_id(scan["model_id"])
    findings = await list_findings(scan_id=scan_id)

    report_id = f"rep-{uuid.uuid4().hex[:8]}"
    now_iso = datetime.now(timezone.utc).isoformat()

    os.makedirs(settings.REPORTS_DIR, exist_ok=True)
    pdf_filename = f"{report_id}.pdf"
    docx_filename = f"{report_id}.docx"
    pdf_path = os.path.join(settings.REPORTS_DIR, pdf_filename)
    docx_path = os.path.join(settings.REPORTS_DIR, docx_filename)

    report_data = {
        "id": report_id,
        "scan_id": scan_id,
        "report_type": report_type,
        "title": title,
        "created_at": now_iso,
        "scan": scan,
        "model": model,
        "findings": findings,
    }

    # Generate physical files
    generate_pdf_report(report_data, pdf_path)
    generate_docx_report(report_data, docx_path)

    doc_record = {
        "id": report_id,
        "scan_id": scan_id,
        "model_id": scan["model_id"],
        "report_type": report_type,
        "title": title,
        "created_at": now_iso,
        "pdf_path": pdf_path,
        "docx_path": docx_path,
        "metadata": {
            "model_name": model.get("name"),
            "risk_score": scan.get("risk_score"),
            "findings_count": len(findings),
        },
    }

    await db["reports"].insert_one(doc_record)
    logger.info(f"Generated report {report_id} (PDF & DOCX) for scan {scan_id}")
    return await db["reports"].find_one({"id": report_id}, {"_id": 0})


async def list_reports() -> List[Dict[str, Any]]:
    """List all compiled reports."""
    db = get_database()
    cursor = db["reports"].find({}, {"_id": 0}).sort("created_at", -1)
    return await cursor.to_list(length=500)


async def get_report_by_id(report_id: str) -> Dict[str, Any]:
    """Retrieve single report metadata."""
    db = get_database()
    report = await db["reports"].find_one({"id": report_id}, {"_id": 0})
    if not report:
        raise ResourceNotFoundException("Report", report_id)
    return report
